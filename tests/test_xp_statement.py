from datetime import datetime
from decimal import Decimal
import io
import json
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from fin2.imports.generic import stage, commit, reject
from fin2.imports.xp_statement import extract
from fin2.imports import xp_reconciliation as xp
from fin2.portfolio.manual_ledger import create
from tests import test_fin1_import as fixtures
from warehouse.database import connect


def workbook(movements=None, holder='ANA TESTE', number='123456', shift=0, future=False, tag=''):
    book = Workbook(); sheet = book.active
    sheet.title = 'Extrato'
    sheet.cell(2, 7, 'Extrato da conta' + tag)
    sheet.cell(3, 7, 'De: 01/01/2025 Até: 31/12/2025')
    sheet.cell(4, 2, holder); sheet.cell(4, 7, 'Conta XP: ' + number)
    sheet.cell(5, 2, 'Saldo total projetado'); sheet.cell(5, 7, 999999)
    for col, value in zip((2,3,4,6,7), ('Movimentação','Liquidação','Lançamento','Valor (R$)','Saldo (R$)')):
        sheet.cell(8+shift, col+shift, value)
    movements = movements if movements is not None else [('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3 S/ 100',10,110)]
    for row, values in enumerate(movements, 9+shift):
        for col, value in zip((2,3,4,6,7), values):
            if col in (2,3): value = datetime.fromisoformat(value)
            sheet.cell(row, col+shift, value)
    row = 9+shift+len(movements)
    sheet.cell(row,2,'Lançamentos futuros')
    if future: sheet.cell(row+1,2,datetime(2026,1,1)); sheet.cell(row+1,6,10000)
    output=io.BytesIO(); book.save(output); book.close(); return output.getvalue()


class XPParserTests(unittest.TestCase):
    def test_review_table_search_sort_and_pagination(self):
        from types import SimpleNamespace
        from fin2.dashboard.xp_statement_views import approval_table
        rows = [dict(source_locator={'row':i}, description='Movimento', amount=str(i),
                     situation='pending', effective_quantity=str(i),
                     effective_application_name='Fundo de teste',
                     decision={'quantity':str(i), 'reason':'Conferir documento'}) for i in range(1,26)]
        selected = {'import_id':'a'*32, 'rows':rows}
        approval_table(SimpleNamespace(GET={'q':'Conferir documento','sort':'table_quantity',
                                           'dir':'desc','page':'3','situacao':'pending',
                                           'por_pagina':'10'}), selected, review=True)
        self.assertEqual(selected['page'].paginator.count, 25)
        self.assertEqual([r['table_quantity'] for r in selected['rows']], ['5','4','3','2','1'])
        self.assertIn('situacao=pending', selected['table_query'])
        self.assertIn('por_pagina=10', selected['table_query'])
        self.assertEqual(selected['page_size'], '10')
        self.assertNotIn('preview=', selected['table_query'])
        selected['rows'] = rows
        # An unknown size falls back to the default without breaking the query.
        approval_table(SimpleNamespace(GET={'por_pagina':'999'}), selected, review=True)
        self.assertIn('por_pagina=50', selected['table_query'])
        selected['rows'] = rows
        # "todos" keeps every row on a single page.
        approval_table(SimpleNamespace(GET={'por_pagina':'todos','page':'2'}), selected, review=True)
        self.assertEqual(len(selected['rows']), 25)
        self.assertEqual(selected['page'].paginator.num_pages, 1)
        selected['rows'] = rows
        approval_table(SimpleNamespace(GET={'situacao':'excluded'}), selected, review=True)
        self.assertIn('situacao=excluded', selected['table_query'])
        selected['rows'] = rows
        approval_table(SimpleNamespace(GET={'situacao':'bogus'}), selected, review=True)
        self.assertNotIn('situacao', selected['table_query'])
        selected['rows'] = rows
        approval_table(SimpleNamespace(GET={'q':'Fundo de teste'}), selected, review=True)
        self.assertEqual(selected['page'].paginator.count, 25)

    def test_approval_search_and_numeric_sort_before_pagination(self):
        from types import SimpleNamespace
        from fin2.dashboard.xp_statement_views import approval_table
        rows = [dict(source_locator={'row':i}, description=f'Movimento {i}',
                     amount=str(i), balance='1234.50', trade_date='2025-01-02',
                     settlement_date='2025-01-06', errors=[]) for i in range(1, 26)]
        selected = {'import_id':'a'*32, 'rows':rows}
        approval_table(SimpleNamespace(GET={'sort':'amount','dir':'desc'}), selected)
        self.assertEqual([r['amount'] for r in selected['rows']], [str(i) for i in range(25,5,-1)])
        selected['rows'] = rows
        approval_table(SimpleNamespace(GET={'q':'Movimento 25'}), selected)
        self.assertEqual(selected['page'].paginator.count, 1)
        self.assertEqual(selected['rows'][0]['source_locator']['row'], 25)
        for term in ('02/01/25', 'R$ 1.234,50'):
            selected['rows'] = rows
            approval_table(SimpleNamespace(GET={'q':term,'sort':'amount','dir':'desc','page':'2'}), selected)
            self.assertEqual(selected['page'].paginator.count, 25)
            self.assertEqual([r['amount'] for r in selected['rows']], ['5','4','3','2','1'])
            self.assertIn('sort=amount&dir=desc', selected['table_query'])

    def test_shifted_headers_settlement_order_futures_and_decimal(self):
        body=workbook([
            ('2025-01-02','2025-01-06','OPERAÇÕES EM BOLSA PR 02/01/2025 NOTA Nº 54321',-20,100),
            ('2025-01-03','2025-01-03','DIVIDENDOS DE CLIENTES TEST3 S/ 100',10,120),
            ('2025-01-03','2025-01-03','DIVIDENDOS DE CLIENTES TEST3 S/ 100',10,110),
        ],shift=2,future=True)
        metadata,rows=extract(body)
        self.assertEqual(len(rows),3)
        self.assertEqual(metadata['opening_balance'],'100')
        self.assertEqual(metadata['closing_balance'],'100')
        self.assertEqual(metadata['balance_breaks'],0)
        self.assertEqual(rows[0].values['note_number'],'54321')
        self.assertNotEqual(rows[1].values['fingerprint'],rows[2].values['fingerprint'])

    def test_broken_balance_formula_unknown_empty_and_invalid(self):
        meta,rows=extract(workbook([
            ('2025-01-03','2025-01-03','DESCRIÇÃO NOVA',1,999),
            ('2025-01-02','2025-01-02','DIVIDENDOS',10,110)]))
        self.assertTrue(meta['errors']); self.assertEqual(rows[0].values['category'],'unknown')
        meta,rows=extract(workbook([('2025-01-02','2025-01-02','DIVIDENDOS','=1+1',110)]))
        self.assertTrue(rows[0].values['errors'])
        with self.assertRaisesRegex(ValueError,'sem movimentos'): extract(workbook([]))
        with self.assertRaisesRegex(ValueError,'inválido'): extract(b'broken')


class XPFlowTests(unittest.TestCase):
    def setUp(self):
        self.fixture=fixtures.ImportTests(); self.fixture.setUp(); self.fixture.run_import()
        self.db=self.fixture.database; self.documents=self.db.parent/'documents'
        with connect(self.db) as db:
            self.account,self.batch=db.execute('select source_record_id,batch_id from portfolio.account').fetchone()
            db.execute("insert into source_record values (?,?, 'db.sqlite3','fin1_moeda',1,?)",['currency',self.batch,json.dumps({'nome':'Real','abrev':'BRL'})])
            for identifier,kind,legacy,payload in (
                ('investor','fin1_titular',1,{'nome':'Ana Teste'}),
                ('investor2','fin1_titular',2,{'nome':'Bruno Teste'}),
                ('asset','fin1_ativo',1,{'nome':'Teste','abrev':'TEST3'}),
                ('other','fin1_conta',2,{'nome':'XP Bruno','moeda_id':1,'titular_id':2}),
                ('digital','fin1_conta',3,{'nome':'Digital Ana','moeda_id':1,'titular_id':1}),
            ):
                db.execute("insert into source_record values (?,?,'db.sqlite3',?,?,?)",[identifier,self.batch,kind,legacy,json.dumps(payload)])
            db.execute("update source_record set payload=json_merge_patch(payload,?) where record_id=?",[json.dumps({'nome':'XP Ana','moeda_id':1,'titular_id':1}),self.account])
            self.application=db.execute('select source_record_id from portfolio.application').fetchone()[0]
            db.execute("update source_record set payload=json_merge_patch(payload,?) where record_id=?",[json.dumps({'nome':'Teste','ativo_id':1}),self.application])

    def tearDown(self): self.fixture.tearDown()

    def test_reupload_documented_statement_requires_load_approval(self):
        import os
        os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
        import django
        django.setup()
        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import Client, override_settings

        body = workbook()
        identifier = self.stage(body)
        xp.document(self.db, identifier)
        with override_settings(WAREHOUSE_PATH=self.db, DOCUMENT_ROOT=self.documents,
                               WRITE_ENABLED=True, ALLOWED_HOSTS=['testserver']):
            client = Client()
            response = client.post('/fin2/importar/', {
                'file': SimpleUploadedFile('extrato.xlsx', body),
                'adapter': 'xp-account-statement', 'account': self.account,
            }, follow=True)
            self.assertEqual(response.status_code, 200)
            html = response.content.decode()
            self.assertIn('<h1>Aprovar carga</h1>', html)
            self.assertIn('DIVIDENDOS DE CLIENTES TEST3', html)
            self.assertIn('Rejeitar arquivo', html)
            self.assertNotIn('Revisão em lote', html)
            self.assertEqual(len(response.redirect_chain), 1)
            from django.urls import reverse
            response = client.post(reverse('xp-statement-update', args=[identifier]),
                                   {'operation': 'document'}, follow=True)
            self.assertEqual(response.status_code, 200)
            self.assertIn('<h1>Revisão do extrato</h1>', response.content.decode())
        with connect(self.db) as db:
            self.assertEqual(db.execute('select count(*) from ledger.manual_event').fetchone()[0], 0)
            # Approving a load records the statement evidence, not a decision log.
            self.assertTrue(db.execute('select documented_at from ledger.xp_statement where import_id=?',
                                       [identifier]).fetchone()[0])
            self.assertEqual(db.execute("""select count(*) from ledger.audit_log
              where entity_type='file_import_decision'""").fetchone()[0], 0)

    def test_individual_review_loads_into_the_list_modal(self):
        import os
        os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
        import django
        django.setup()
        from django.test import Client, override_settings
        from django.urls import reverse

        identifier = self.stage(workbook())
        xp.document(self.db, identifier)
        with override_settings(WAREHOUSE_PATH=self.db, DOCUMENT_ROOT=self.documents,
                               WRITE_ENABLED=True, ALLOWED_HOSTS=['testserver']):
            client = Client()
            listing = client.get(reverse('xp-statement-review', args=[identifier])).content.decode()
            # The list carries the modal shell and one trigger per movement.
            self.assertIn('data-review-dialog', listing)
            self.assertIn('data-review-modal', listing)
            self.assertNotIn('id="individual-review"', listing)
            focused = client.get(reverse('xp-statement-review', args=[identifier]),
                                 {'review_line': '1'}).content.decode()
            # The focused route stays the fallback and the modal source.
            self.assertIn('id="individual-review"', focused)
            self.assertNotIn('data-review-dialog', focused)
            # The counterparty is always the result account, not a free choice.
            self.assertIn('<input type="hidden" name="counterparty" value="RESULTADO">', focused)
            self.assertNotIn('<select name="counterparty"', focused)
            # The category is editable; the event type is derived from it.
            self.assertIn('<select name="category">', focused)
            self.assertNotIn('name="event_type"', focused)
            # The modal heading shows the same line the list trigger points at.
            import re
            self.assertEqual(re.search(r'data-review-line="(\d+)"', listing).group(1),
                             re.search(r'Revisão da Linha (\d+)', focused).group(1))

    def test_individual_review_infers_the_action_without_a_decision_field(self):
        import os
        os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
        import django
        django.setup()
        from django.test import Client, override_settings
        from django.urls import reverse

        identifier = self.stage(workbook())
        xp.document(self.db, identifier)
        with override_settings(WAREHOUSE_PATH=self.db, DOCUMENT_ROOT=self.documents,
                               WRITE_ENABLED=True, ALLOWED_HOSTS=['testserver']):
            client = Client()
            url = reverse('xp-statement-update', args=[identifier])
            # No decision field: an entry-less save becomes a new entry, and the
            # ledger event follows from the confirmed category.
            client.post(url, {'line_number': '1', 'reason': 'Provento conferido',
                              'category': 'dividend', 'application_record': self.application})
            with connect(self.db) as db:
                payload = json.loads(db.execute('select decision from ledger.xp_statement_line'
                                                ' where import_id=? and line_number=1', [identifier]).fetchone()[0])
            self.assertEqual(payload['action'], 'new')
            self.assertEqual(payload['category'], 'dividend')
            self.assertEqual(payload['event_type'], 'income')
            # "Alterar sugestão" replaces the identified application: an unknown
            # override is refused even though the suggestion itself is valid.
            client.post(url, {'line_number': '1', 'reason': 'Provento conferido',
                              'category': 'dividend', 'application_record': self.application,
                              'application_record_override': 'inexistente'})
            with connect(self.db) as db:
                payload = json.loads(db.execute('select decision from ledger.xp_statement_line'
                                                ' where import_id=? and line_number=1', [identifier]).fetchone()[0])
            self.assertEqual(payload['application_record'], self.application)
            # The dedicated button excludes the line without touching other fields.
            client.post(url, {'line_number': '1', 'action': 'excluded', 'reason': 'Duplicado no extrato'})
            with connect(self.db) as db:
                payload = json.loads(db.execute('select decision from ledger.xp_statement_line'
                                                ' where import_id=? and line_number=1', [identifier]).fetchone()[0])
            self.assertEqual(payload['action'], 'excluded')

    def test_reviewed_category_selects_the_event_type(self):
        identifier = self.stage(workbook())
        # The extracted movement is a dividend; reviewing it as a transfer turns
        # it into an external contribution because the counterparty is the result
        # account.
        xp.review(self.db, identifier, 1, {'action': 'new', 'category': 'transfer',
                                           'reason': 'Aporte externo conferido',
                                           'distinct_confirmed': True})
        with connect(self.db) as db:
            payload = json.loads(db.execute('select decision from ledger.xp_statement_line'
                                            ' where import_id=? and line_number=1', [identifier]).fetchone()[0])
        self.assertEqual(payload['category'], 'transfer')
        self.assertEqual(payload['event_type'], 'deposit')
        self.assertIsNone(payload['counterparty'])

    def test_commit_button_only_with_ready_lines_and_error_keeps_step_three(self):
        from urllib.parse import unquote_plus
        import os
        os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
        import django
        django.setup()
        from django.test import Client, override_settings
        from django.urls import reverse

        # Brokerage needs a link to an existing record, so nothing is ready.
        nothing_ready = self.stage(workbook([
            ('2025-01-02','2025-01-02','OPERAÇÕES EM BOLSA PR 02/01/2025 NOTA Nº 123',90,190)]))
        partly_ready = self.stage(workbook([
            ('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,110),
            ('2025-01-03','2025-01-03','OPERAÇÕES EM BOLSA PR 03/01/2025 NOTA Nº 124',5,115)]))
        all_ready = self.stage(workbook([('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,110)]))
        for identifier in (nothing_ready, partly_ready, all_ready):
            xp.document(self.db, identifier)
        with override_settings(WAREHOUSE_PATH=self.db, DOCUMENT_ROOT=self.documents,
                               WRITE_ENABLED=True, ALLOWED_HOSTS=['testserver']):
            client = Client()
            empty = reverse('xp-statement-review', args=[nothing_ready])
            self.assertIn('class="button-commit" disabled>', client.get(empty).content.decode())
            review = reverse('xp-statement-review', args=[partly_ready])
            html = client.get(review).content.decode()
            self.assertIn('class="button-commit">', html)
            # Pending lines still block step 4, and the reviewer stays in step 3.
            response = client.post(reverse('file-import-commit', args=[partly_ready]))
            self.assertEqual(response.status_code, 302)
            self.assertTrue(response['Location'].startswith(review + '?'))
            self.assertIn('Resolva todas as linhas', unquote_plus(response['Location']))
            with connect(self.db) as db:
                self.assertEqual(db.execute('select count(*) from ledger.manual_event').fetchone()[0], 0)
            # With every line resolved the load concludes and step 4 lists what was created.
            response = client.post(reverse('file-import-commit', args=[all_ready]))
            self.assertEqual(response.status_code, 302)
            self.assertEqual(response['Location'], '/fin2/importar/?preview='+all_ready)
            page = client.get(response['Location']).content.decode()
            self.assertIn('Importação de extrato concluída', page)
            self.assertIn('DIVIDENDOS DE CLIENTES TEST3', page)

    def test_rejected_import_is_discarded_and_returns_to_step_one(self):
        import os
        os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
        import django
        django.setup()
        from django.test import Client, override_settings
        from django.urls import reverse

        identifier = self.stage(workbook())
        with override_settings(WAREHOUSE_PATH=self.db, DOCUMENT_ROOT=self.documents,
                               WRITE_ENABLED=True, ALLOWED_HOSTS=['testserver']):
            client = Client()
            response = client.post(reverse('file-import-reject', args=[identifier]))
            # Back to the upload step: no preview is selected any more.
            self.assertEqual(response.status_code, 302)
            self.assertEqual(response['Location'], '/fin2/importar/')
            with connect(self.db) as db:
                self.assertIsNone(db.execute('select 1 from ledger.file_import where import_id=?',
                                             [identifier]).fetchone())
                self.assertIsNone(db.execute('select 1 from ledger.xp_statement where import_id=?',
                                             [identifier]).fetchone())
                self.assertEqual(db.execute('select count(*) from ledger.xp_statement_line').fetchone()[0], 0)
                self.assertEqual(db.execute("""select count(*) from ledger.audit_log
                  where entity_type='file_import_decision'""").fetchone()[0], 0)
            self.assertEqual(client.get('/fin2/importar/', {'preview': identifier}).status_code, 404)
            self.assertEqual(client.get(reverse('xp-statement-review', args=[identifier])).status_code, 404)

    def test_document_records_evidence_without_a_decision_log(self):
        identifier = self.stage(workbook())
        xp.document(self.db, identifier)
        with connect(self.db) as db:
            self.assertTrue(db.execute('select documented_at from ledger.xp_statement where import_id=?',
                                       [identifier]).fetchone()[0])
            self.assertEqual(db.execute('select count(*) from ledger.statement_balance_observation').fetchone()[0], 1)
            self.assertEqual(db.execute("""select count(*) from ledger.audit_log
              where entity_type='file_import_decision'""").fetchone()[0], 0)
        reject(self.db, self.documents, identifier)
        with connect(self.db) as db:
            self.assertIsNone(db.execute('select 1 from ledger.file_import where import_id=?', [identifier]).fetchone())
            self.assertIsNone(db.execute('select 1 from ledger.xp_statement where import_id=?', [identifier]).fetchone())
            self.assertEqual(db.execute('select count(*) from ledger.xp_statement_line').fetchone()[0], 0)
            self.assertEqual(db.execute('select count(*) from ledger.statement_balance_observation').fetchone()[0], 0)

    def test_excluded_line_is_ready_to_finish_and_never_reaches_the_ledger(self):
        identifier = self.stage(workbook([
            ('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3 S/ 100',10,110),
            ('2025-01-03','2025-01-03','MOVIMENTO SEM DOCUMENTO DE SUPORTE',5,115)]))
        xp.review(self.db, identifier, 1, self.decision())
        xp.review(self.db, identifier, 2, {'action':'excluded','reason':'Linha sem documento de suporte'})
        with connect(self.db) as db:
            detail = xp.detail(db, identifier)
            self.assertEqual({r['row_number']: r['state'] for r in detail['rows']}, {1:'ready', 2:'excluded'})
            self.assertEqual(detail['states'], {'ready':1, 'pending':0, 'excluded':1})
            # Nothing is left pending, so the load can be posted.
            self.assertTrue(detail['can_confirm'])
            self.assertTrue(detail['can_commit'])
            history = db.execute("""select json_extract_string(payload,'$.action') from ledger.xp_statement_decision
              where import_id=? order by line_number""", [identifier]).fetchall()
            self.assertEqual([a for (a,) in history], ['new', 'excluded'])
        self.assertEqual(commit(self.db, identifier), 1)
        with connect(self.db) as db:
            self.assertEqual(db.execute('select count(*) from ledger.manual_event').fetchone()[0], 1)
            # Step 4 discards the review staging but keeps link and claim.
            for table in ('xp_statement', 'xp_statement_line', 'xp_statement_decision'):
                self.assertEqual(db.execute(f'select count(*) from ledger.{table}').fetchone()[0], 0)
            self.assertEqual(db.execute('select count(*) from ledger.xp_statement_link where line_number=2').fetchone()[0], 0)
            self.assertEqual(db.execute('select count(*) from ledger.xp_statement_claim where line_number=2').fetchone()[0], 0)
            self.assertEqual(db.execute('select count(*) from ledger.xp_statement_claim').fetchone()[0], 1)
            created = xp.created_entries(db, identifier)
        self.assertEqual([e['event_type'] for e in created], ['income'])
        self.assertEqual(created[0]['description'], 'DIVIDENDOS DE CLIENTES TEST3 S/ 100')
        import os
        os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
        import django
        django.setup()
        from django.test import Client, override_settings
        with override_settings(WAREHOUSE_PATH=self.db, DOCUMENT_ROOT=self.documents,
                               WRITE_ENABLED=True, ALLOWED_HOSTS=['testserver']):
            response = Client().get('/fin2/importar/', {'preview': identifier})
            self.assertEqual(response.status_code, 200)
            html = response.content.decode()
            # Step 4 reports the load and lists the entries it generated.
            self.assertIn('Importação de extrato concluída. Novos lançamentos gerados no ledger: 1.', html)
            self.assertIn('DIVIDENDOS DE CLIENTES TEST3 S/ 100', html)
            self.assertNotIn('Histórico de revisões', html)

    def test_overlap_check_survives_the_concluded_load(self):
        first = self.stage(workbook([('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,110)]))
        xp.review(self.db, first, 1, self.decision())
        self.assertEqual(commit(self.db, first), 1)
        # The concluded load kept no lines, only the claim; the overlap check
        # reads it there and must still flag a reissued statement with a new value.
        second = self.stage(workbook([('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',7,107)]))
        with self.assertRaisesRegex(ValueError, 'sobreposto'):
            xp.review(self.db, second, 1, self.decision())

    def test_entry_linked_twice_is_refused_after_the_load_concludes(self):
        entry = create(self.db, account_record=self.account, event_type='deposit',
                       settlement_date='2025-01-02', currency='BRL', amount=10, description='Aporte conferido')
        first = self.stage(workbook([('2025-01-02','2025-01-02','APORTE RECEBIDO',10,110)]))
        xp.review(self.db, first, 1, self.decision(action='link', entry_ids=[entry]))
        self.assertEqual(commit(self.db, first), 0)
        # The concluded load kept the link, so the same movement cannot be
        # reconciled again from a divergent line of another statement.
        second = self.stage(workbook([('2025-01-02','2025-01-02','APORTE RECEBIDO OUTRA CONSULTA',10,110)]))
        with self.assertRaisesRegex(ValueError, 'outra linha'):
            xp.review(self.db, second, 1, self.decision(action='link', entry_ids=[entry]))

    def stage(self,body=None,account=None,**options):
        return stage(self.db,self.documents,'extrato.xlsx',body or workbook(),options=dict(account_record=account or self.account,confirm_identity=True,identity_reason='Número e titular conferidos',**options))[0]

    def decision(self,**changes):
        return dict(dict(action='new',event_type='income',application_record=self.application,reason='Provento conferido'),**changes)

    def test_evidence_then_income_and_reexport_no_duplicate(self):
        body=workbook(); identifier=self.stage(body)
        self.assertEqual(self.stage(body),identifier)
        xp.document(self.db,identifier); xp.document(self.db,identifier)
        with connect(self.db) as db:
            self.assertEqual(db.execute('select count(*) from ledger.manual_event').fetchone()[0],0)
            self.assertTrue(db.execute('select opening_inferred from ledger.statement_balance_observation').fetchone()[0])
            self.assertEqual(xp.detail(db,identifier)['counts']['pending'],1)
        xp.review(self.db,identifier,1,self.decision())
        self.assertEqual(commit(self.db,identifier),1)
        second=self.stage(workbook(tag=' reexportado'))
        self.assertEqual(commit(self.db,second),0)
        with connect(self.db) as db:
            self.assertEqual(db.execute('select count(*) from ledger.manual_event').fetchone()[0],1)
            self.assertEqual(db.execute('select count(*) from ledger.xp_statement_link').fetchone()[0],2)
            self.assertEqual(db.execute('select count(*) from ledger.purchase_funding').fetchone()[0],0)
        with self.assertRaises(ValueError): commit(self.db,identifier)

    def test_multiple_holders_and_identity_guard(self):
        self.stage()
        second=self.stage(workbook(holder='BRUNO TESTE',number='777888'),account='other')
        self.assertTrue(second)
        with self.assertRaisesRegex(ValueError,'diverge'): self.stage(workbook(tag='errado'),account='other')
        with self.assertRaisesRegex(ValueError,'Titular'): self.stage(workbook(holder='PESSOA ERRADA',number='999'),account='digital')
        with self.assertRaisesRegex(ValueError,'Confirme'):
            stage(self.db,self.documents,'x.xlsx',workbook(holder='ANA TESTE',number='111'),options={'account_record':'digital'})

    def test_pending_rejection_and_atomic_rollback(self):
        identifier=self.stage(workbook([
            ('2025-01-03','2025-01-03','DIVIDENDOS DE CLIENTES TEST3',20,130),
            ('2025-01-02','2025-01-02','OPERAÇÕES EM BOLSA PR 02/01/2025 NOTA Nº 123',10,110)]))
        xp.review(self.db,identifier,1,self.decision())
        # Brokerage has no ledger event of its own, so the line stays pending and
        # blocks the confirmation until it is linked or excluded.
        with self.assertRaisesRegex(ValueError,'todas'): commit(self.db,identifier)
        xp.review(self.db,identifier,2,self.decision(action='excluded',reason='Sem nota correspondente'))
        original=xp._create
        def fail_first(db,item,row,decision):
            if row['row_number']==1: raise ValueError('falha simulada')
            return original(db,item,row,decision)
        with patch.object(xp,'_create',side_effect=fail_first):
            with self.assertRaisesRegex(ValueError,'simulada'): commit(self.db,identifier)
        with connect(self.db) as db:
            self.assertEqual(db.execute('select count(*) from ledger.manual_event').fetchone()[0],0)
            self.assertEqual(db.execute('select count(*) from ledger.xp_statement_claim').fetchone()[0],0)
        reject(self.db,self.documents,identifier)
        with self.assertRaises(ValueError): commit(self.db,identifier)

    def test_brokerage_many_to_one_and_no_invented_trade(self):
        body=workbook([('2025-01-02','2025-01-02','OPERAÇÕES EM BOLSA PR 02/01/2025 NOTA Nº 123',90,190)])
        identifier=self.stage(body)
        with self.assertRaisesRegex(ValueError,'Categoria'): xp.review(self.db,identifier,1,self.decision())
        events=[create(self.db,account_record=self.account,event_type=kind,settlement_date='2025-01-02',currency='BRL',amount=amount,description=description)
                for kind,amount,description in [('deposit',100,'Nota 123'),('fee',-10,'Despesas nota 123')]]
        xp.review(self.db,identifier,1,self.decision(action='link',entry_ids=events))
        self.assertEqual(commit(self.db,identifier),0)
        with connect(self.db) as db: self.assertEqual(db.execute('select count(*) from ledger.xp_statement_link').fetchone()[0],2)

    def test_duplicate_previews_revalidate_new_existing_event(self):
        first=self.stage(); second=self.stage(workbook(tag='outra consulta'))
        for identifier in (first,second): xp.review(self.db,identifier,1,self.decision())
        self.assertEqual(commit(self.db,first),1); self.assertEqual(commit(self.db,second),0)
        third=self.stage(workbook([('2025-01-04','2025-01-04','DIVIDENDOS DE CLIENTES TEST3',10,110)]))
        xp.review(self.db,third,1,self.decision())
        create(self.db,account_record=self.account,event_type='income',settlement_date='2025-01-04',currency='BRL',amount=10,description='Lançado depois da prévia')
        with self.assertRaisesRegex(ValueError,'mesma data'): commit(self.db,third)

    def test_own_transfer_and_wrong_holder(self):
        identifier=self.stage(workbook([('2025-01-02','2025-01-02','Transferência enviada para a conta digital',-10,90)]))
        with self.assertRaisesRegex(ValueError,'mesmo titular'):
            xp.review(self.db,identifier,1,self.decision(event_type='transfer',counterparty='other'))
        xp.review(self.db,identifier,1,self.decision(event_type='transfer',counterparty='digital'))
        self.assertEqual(commit(self.db,identifier),2)
        with connect(self.db) as db:
            self.assertEqual(db.execute('select sum(amount) from ledger.manual_event').fetchone()[0],0)
            self.assertEqual(db.execute("select count(*) from ledger.current_account_entry where entry_id in (select event_id from ledger.manual_event)").fetchone()[0],0)

    def test_repeated_identical_descriptions_are_distinct(self):
        identifier=self.stage(workbook([
            ('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,120),
            ('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,110)]))
        for line in (1,2): xp.review(self.db,identifier,line,self.decision())
        self.assertEqual(commit(self.db,identifier),2)
        with connect(self.db) as db: self.assertEqual(db.execute('select count(*) from ledger.xp_statement_claim').fetchone()[0],2)

    def test_redemption_tax_requires_pair_and_preserves_two_amounts(self):
        identifier=self.stage(workbook([
            ('2025-01-02','2025-01-02','IRRF S/RESGATE FUNDOS Teste',-5,195),
            ('2025-01-02','2025-01-02','RESGATE Teste',100,200)]))
        with self.assertRaisesRegex(ValueError,'Vincule o imposto'):
            xp.review(self.db,identifier,1,self.decision(event_type='tax'))
        with self.assertRaisesRegex(ValueError,'quantidade'):
            xp.review(self.db,identifier,2,self.decision(event_type='redemption'))
        with connect(self.db) as db:
            document=db.execute('select document_id from source_document where original_filename=?',['note.pdf']).fetchone()[0]
        xp.review(self.db,identifier,1,self.decision(event_type='tax',related_line='2'))
        xp.review(self.db,identifier,2,self.decision(event_type='redemption',quantity='2',document_id=document))
        self.assertEqual(commit(self.db,identifier),2)
        with connect(self.db) as db:
            self.assertEqual(db.execute('select sum(amount) from ledger.manual_event').fetchone()[0],Decimal('95'))
            self.assertEqual(db.execute("select quantity from ledger.manual_event where event_type='redemption'").fetchone()[0],Decimal('2'))

    def test_one_existing_event_cannot_cover_two_rows(self):
        identifier=self.stage(workbook([
            ('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,120),
            ('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,110)]))
        event=create(self.db,account_record=self.account,event_type='income',settlement_date='2025-01-02',currency='BRL',amount=10,description='Provento')
        for line in (1,2): xp.review(self.db,identifier,line,self.decision(action='link',entry_ids=[event]))
        with self.assertRaisesRegex(ValueError,'duas linhas'): commit(self.db,identifier)

    def test_short_holder_names_are_explicitly_associated(self):
        with connect(self.db) as db:
            db.execute("update source_record set payload=? where record_id='investor'",[json.dumps({'nome':'Ana'})])
        identifier=self.stage()
        self.assertTrue(identifier)
        with connect(self.db) as db:
            self.assertEqual(db.execute('select holder from ledger.xp_account_binding').fetchone()[0],'ANA TESTE')

    def test_identification_marks_ready_lines_and_commit_records_them(self):
        identifier=self.stage(workbook([
            ('2025-01-04','2025-01-04','DIVIDENDOS DE CLIENTES TEST3',20,135),
            ('2025-01-03','2025-01-03','DIVIDENDOS DE CLIENTES TEST3',5,115),
            ('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,110)]))
        create(self.db,account_record=self.account,event_type='income',settlement_date='2025-01-03',currency='BRL',amount=5,description='DIVIDENDOS  DE CLIENTES TEST3')
        with connect(self.db) as db: item=xp.detail(db,identifier)
        self.assertEqual(item['states'],{'ready':3,'pending':0,'excluded':0})
        self.assertEqual(item['rows'][1]['identified']['action'],'link')
        # The identification is a preview: nothing is written before confirming.
        with connect(self.db) as db:
            self.assertEqual(db.execute('select count(*) from ledger.xp_statement_decision').fetchone()[0],0)
            self.assertEqual(db.execute('select count(*) from ledger.manual_event').fetchone()[0],1)
        self.assertEqual(commit(self.db,identifier),2)

    def test_manual_decision_wins_over_the_identification(self):
        identifier=self.stage(workbook([('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,110)]))
        create(self.db,account_record=self.account,event_type='income',settlement_date='2025-01-02',currency='BRL',amount=10,description='DIVIDENDOS DE CLIENTES TEST3')
        with connect(self.db) as db: item=xp.detail(db,identifier)
        self.assertEqual(item['rows'][0]['identified']['action'],'link')
        xp.review(self.db,identifier,1,self.decision(action='excluded',reason='Duplicado do extrato'))
        with connect(self.db) as db:
            item=xp.detail(db,identifier)
            self.assertIsNone(item['rows'][0]['identified'])
            self.assertEqual(item['rows'][0]['state'],'excluded')

    def test_ambiguous_candidates_stay_pending_with_a_hint(self):
        identifier=self.stage(workbook([('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,110)]))
        for _ in range(2):
            create(self.db,account_record=self.account,event_type='income',settlement_date='2025-01-02',currency='BRL',amount=10,description='DIVIDENDOS DE CLIENTES TEST3')
        with connect(self.db) as db: item=xp.detail(db,identifier)
        self.assertEqual(item['rows'][0]['state'],'pending')
        self.assertEqual(item['rows'][0]['identification_missing'],['Escolha entre os movimentos existentes'])

    def test_commit_rolls_back_every_line_on_error(self):
        identifier=self.stage(workbook([
            ('2025-01-03','2025-01-03','DIVIDENDOS DE CLIENTES TEST3',20,130),
            ('2025-01-02','2025-01-02','DIVIDENDOS DE CLIENTES TEST3',10,110)]))
        original=xp._create
        def fail_second(db,item,row,decision):
            if row['row_number']==2: raise ValueError('falha simulada na confirmação')
            return original(db,item,row,decision)
        with patch.object(xp,'_create',side_effect=fail_second):
            with self.assertRaisesRegex(ValueError,'simulada'): commit(self.db,identifier)
        with connect(self.db) as db:
            self.assertEqual(db.execute('select count(*) from ledger.manual_event').fetchone()[0],0)
            self.assertEqual(db.execute('select count(*) from ledger.xp_statement_claim').fetchone()[0],0)


    def test_resultado_counterparty_becomes_deposit(self):
        body = workbook(movements=[('2025-01-02', '2025-01-02', 'Transferência', 150.00, 150.00)])
        identifier = self.stage(body)
        from fin2.imports import xp_reconciliation as xp
        from warehouse.database import connect
        xp.document(self.db, identifier)
        with connect(self.db) as db:
            item_before = xp._load(db, identifier, True)
            row_num = item_before['rows'][0]['row_number']

        # Review as transfer but with RESULTADO
        xp.review(self.db, identifier, row_num, self.decision(
            action='new', event_type='transfer', counterparty='RESULTADO', reason='From P&L'
        ))

        # Check if it was mutated
        with connect(self.db) as db:
            item = xp._load(db, identifier, True)
            self.assertEqual(item['rows'][0]['decision']['event_type'], 'deposit')
            self.assertIsNone(item['rows'][0]['decision']['counterparty'])

        # Commit should create a deposit, not a transfer
        xp.commit(self.db, identifier)
        with connect(self.db) as db:
            events = db.execute('select event_type, amount, transfer_id from ledger.manual_event').fetchall()
            self.assertEqual(len(events), 1)
            self.assertEqual(events[0][0], 'deposit')
            self.assertIsNone(events[0][2]) # transfer_id should be None

    def test_resultado_counterparty_becomes_withdrawal(self):
        body = workbook(movements=[('2025-01-02', '2025-01-02', 'Transferência', -150.00, -150.00)])
        identifier = self.stage(body)
        from fin2.imports import xp_reconciliation as xp
        from warehouse.database import connect
        xp.document(self.db, identifier)
        with connect(self.db) as db:
            item_before = xp._load(db, identifier, True)
            row_num = item_before['rows'][0]['row_number']

        xp.review(self.db, identifier, row_num, self.decision(
            action='new', event_type='transfer', counterparty='RESULTADO', reason='To P&L'
        ))

        with connect(self.db) as db:
            item = xp._load(db, identifier, True)
            self.assertEqual(item['rows'][0]['decision']['event_type'], 'withdrawal')

        xp.commit(self.db, identifier)
        with connect(self.db) as db:
            events = db.execute('select event_type, amount, transfer_id from ledger.manual_event').fetchall()
            self.assertEqual(len(events), 1)
            self.assertEqual(events[0][0], 'withdrawal')

class IdentifyApplicationTests(unittest.TestCase):
    """Unit tests for the identify_application() cascade."""

    def _app(self, source_record_id, name, symbol=None, asset_name=None,
             cnpj=None, issuer=None, asset_aliases=None, app_aliases=None):
        return {
            'source_record_id': source_record_id, 'name': name,
            'symbol': symbol, 'asset_name': asset_name,
            'cnpj': cnpj, 'issuer': issuer,
            'asset_aliases': asset_aliases, 'app_aliases': app_aliases,
        }

    def _row(self, description, symbol=None, category='income'):
        return {'description': description, 'symbol': symbol, 'category': category}

    # ---- Step 1: exact symbol -----------------------------------------------

    def test_ticker_exact_unique(self):
        apps = [self._app('a1', 'PETR4 XP', symbol='PETR4')]
        row = self._row('JUROS S/ CAPITAL PROPRIO S/100 PETR4 - PN', symbol='PETR4')
        app_id, method = xp.identify_application(row, apps)
        self.assertEqual(app_id, 'a1')
        self.assertEqual(method, 'exact_symbol')

    def test_ticker_ambiguous_two_apps(self):
        apps = [self._app('a1', 'PETR4 A', symbol='PETR4'),
                self._app('a2', 'PETR4 B', symbol='PETR4')]
        row = self._row('DIVIDENDOS PETR4', symbol='PETR4')
        app_id, method = xp.identify_application(row, apps)
        self.assertIsNone(app_id)

    # ---- Step 2: CNPJ -------------------------------------------------------

    def test_cnpj_formatted(self):
        apps = [self._app('a1', 'Fundo A', cnpj='12345678000190')]
        row = self._row('RENDIMENTOS 12.345.678/0001-90 CDB')
        app_id, method = xp.identify_application(row, apps)
        self.assertEqual(app_id, 'a1')
        self.assertEqual(method, 'cnpj')

    def test_cnpj_no_match(self):
        apps = [self._app('a1', 'Fundo A', cnpj='00000000000000')]
        row = self._row('RENDIMENTOS 12.345.678/0001-90 CDB')
        app_id, _ = xp.identify_application(row, apps)
        self.assertIsNone(app_id)

    # ---- Step 3: aliases ----------------------------------------------------

    def test_asset_alias_match(self):
        apps = [self._app('a1', 'Trend Ações FIA', asset_aliases='TREND ACOES FI|TREND ACOES')]
        row = self._row('RESGATE TREND ACOES FI')
        app_id, method = xp.identify_application(row, apps)
        self.assertEqual(app_id, 'a1')
        self.assertEqual(method, 'alias')

    def test_app_alias_match(self):
        apps = [self._app('a1', 'CDB Inter', app_aliases='CDB BANCO INTER')]
        row = self._row('RENDIMENTOS CDB BANCO INTER 2025')
        app_id, method = xp.identify_application(row, apps)
        self.assertEqual(app_id, 'a1')
        self.assertEqual(method, 'alias')

    def test_alias_ambiguous(self):
        apps = [self._app('a1', 'Fundo X', asset_aliases='TREND'),
                self._app('a2', 'Fundo Y', app_aliases='TREND')]
        row = self._row('RESGATE TREND ACOES FI')
        app_id, _ = xp.identify_application(row, apps)
        self.assertIsNone(app_id)

    # ---- Step 4: asset name -------------------------------------------------

    def test_asset_name_match(self):
        apps = [self._app('a1', 'Trend', asset_name='TREND ACOES FIA')]
        row = self._row('RESGATE TREND ACOES FIA')
        app_id, method = xp.identify_application(row, apps)
        self.assertEqual(app_id, 'a1')
        self.assertEqual(method, 'asset_name')

    # ---- Step 5: application name -------------------------------------------

    def test_app_name_match(self):
        apps = [self._app('a1', 'CICLO DE VIDA 2030')]
        row = self._row('APLICACAO PREVIDENCIA CICLO DE VIDA 2030', category='pension')
        app_id, method = xp.identify_application(row, apps)
        self.assertEqual(app_id, 'a1')
        self.assertEqual(method, 'app_name')

    # ---- Step 6: issuer -----------------------------------------------------

    def test_issuer_match(self):
        apps = [self._app('a1', 'CDB Inter', issuer='BANCO INTER')]
        row = self._row('RENDIMENTOS CDB BANCO INTER')
        app_id, method = xp.identify_application(row, apps)
        self.assertEqual(app_id, 'a1')
        self.assertEqual(method, 'issuer')

    # ---- No match / empty ---------------------------------------------------

    def test_no_apps(self):
        app_id, method = xp.identify_application(self._row('TRANSFERENCIA'), [])
        self.assertIsNone(app_id)
        self.assertIsNone(method)

    def test_generic_transfer_no_match(self):
        apps = [self._app('a1', 'PETR4 XP', symbol='PETR4')]
        app_id, method = xp.identify_application(
            self._row('TRANSFERENCIA DA CONTA DIGITAL', category='transfer'), apps)
        self.assertIsNone(app_id)


class ValidateAliasesTests(unittest.TestCase):
    """Unit tests for catalog._validate_aliases()."""

    def test_valid_aliases(self):
        from fin2.portfolio.catalog import _validate_aliases
        payload = {'statement_aliases': 'Alias A | Alias B | Alias C'}
        _validate_aliases(payload)
        self.assertEqual(payload['statement_aliases'], 'Alias A|Alias B|Alias C')

    def test_empty_aliases_becomes_none(self):
        from fin2.portfolio.catalog import _validate_aliases
        payload = {'statement_aliases': ''}
        _validate_aliases(payload)
        self.assertIsNone(payload['statement_aliases'])

    def test_missing_field_becomes_none(self):
        from fin2.portfolio.catalog import _validate_aliases
        payload = {}
        _validate_aliases(payload)
        self.assertIsNone(payload['statement_aliases'])

    def test_too_many_aliases(self):
        from fin2.portfolio.catalog import _validate_aliases
        payload = {'statement_aliases': '|'.join([f'a{i}' for i in range(21)])}
        with self.assertRaises(ValueError):
            _validate_aliases(payload)

    def test_alias_too_long(self):
        from fin2.portfolio.catalog import _validate_aliases
        payload = {'statement_aliases': 'A' * 101}
        with self.assertRaises(ValueError):
            _validate_aliases(payload)
