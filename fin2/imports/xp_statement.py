"""XP account statement extraction. No workbook content is executable."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import re
import unicodedata
import zipfile

from openpyxl import load_workbook
from fin2.imports.base import SourceRow
from fin2.imports.registry import register

MAX_ROWS = 20000


def plain(value):
    text = unicodedata.normalize('NFKD', str(value or ''))
    return ' '.join(''.join(c for c in text if not unicodedata.combining(c)).upper().split())


def money(value):
    if isinstance(value, str):
        value = value.strip().replace('R$', '').replace(' ', '')
        if ',' in value:
            value = value.replace('.', '').replace(',', '.')
    try:
        result = Decimal(str(value))
        if not result.is_finite() or result != result.quantize(Decimal('.01')):
            raise ValueError()
        return result
    except (InvalidOperation, ValueError):
        raise ValueError('Valor monetário inválido') from None


def day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError('Data inválida')


def category(description):
    text = plain(description)
    for prefix, kind in (
        ('JUROS S/ CAPITAL', 'jcp'), ('DIVIDENDOS', 'dividend'),
        ('RENDIMENTOS', 'income'), ('OPERACOES EM BOLSA', 'brokerage'),
        ('TRANSFERENCIA', 'transfer'), ('TED ', 'transfer'),
        ('IRRF S/RESGATE', 'redemption_tax'), ('RESGATE ', 'redemption'),
        ('APLICACAO PREVIDENCIA', 'pension'),
    ):
        if text.startswith(prefix):
            return kind
    return 'unknown'


def extract(body):
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            files = archive.infolist()
            if len(files) > 1000 or sum(f.file_size for f in files) > 30 * 1024 * 1024:
                raise ValueError('XLSX excede o limite descompactado')
            if any('externalLinks/' in f.filename or 'vbaProject' in f.filename for f in files):
                raise ValueError('XLSX com vínculos externos ou macros não é aceito')
        book = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError('XLSX inválido') from exc
    metadata = {'currency': 'BRL', 'opening_inferred': True}
    rows = []
    try:
        if len(book.worksheets) > 10:
            raise ValueError('XLSX com excesso de abas')
        for sheet in book:
            if sheet.max_row > MAX_ROWS or sheet.max_column > 100:
                raise ValueError('XLSX excede o limite de linhas/colunas')
            headers = None
            future = False
            for number, cells in enumerate(sheet.iter_rows(), 1):
                if number > MAX_ROWS or len(cells) > 100:
                    raise ValueError('XLSX excede o limite de linhas/colunas')
                values = [c.value for c in cells]
                texts = [plain(v) for v in values]
                joined = ' '.join(str(v) for v in values if v is not None)
                if 'CONTA XP:' in plain(joined):
                    match = re.search(r'Conta XP:\s*([\d.\-]+)', joined, re.I)
                    if not match:
                        raise ValueError('Número de conta XP inválido')
                    account_number = re.sub(r'\D', '', match[1]).lstrip('0') or '0'
                    if metadata.get('account_number', account_number) != account_number:
                        raise ValueError('Arquivo contém contas XP diferentes')
                    metadata['account_number'] = account_number
                    holder = next((str(v).strip() for v in values if v and 'CONTA XP:' not in plain(v)), '')
                    metadata['holder'] = holder
                period = re.search(r'De:\s*(\d{2}/\d{2}/\d{4})\s+Até:\s*(\d{2}/\d{2}/\d{4})', joined, re.I)
                if period:
                    metadata.update(period_start=day(period[1]), period_end=day(period[2]))
                if 'DATA DA CONSULTA:' in plain(joined):
                    metadata['queried_at'] = joined.strip()
                if 'SALDO TOTAL PROJETADO' in texts:
                    metadata['projected_balance'] = next((str(v) for v in reversed(values) if isinstance(v, (int, float))), None)
                if 'LANCAMENTOS FUTUROS' in texts:
                    future = True
                    headers = None
                    continue
                if future:
                    continue
                needed = ('MOVIMENTACAO', 'LIQUIDACAO', 'LANCAMENTO', 'VALOR (R$)', 'SALDO (R$)')
                if all(k in texts for k in needed):
                    headers = [texts.index(k) for k in needed]
                    continue
                if headers is None or not any(v is not None and str(v).strip() for v in values):
                    continue
                if any(t.startswith(('IMPORTANTE:', 'EXTRATO PARA SIMPLES', 'PARA RECLAMACOES')) for t in texts):
                    headers = None
                    continue
                selected = [values[i] for i in headers]
                if plain(selected[2]) == 'NAO HA LANCAMENTOS PARA O PERIODO':
                    continue
                errors = []
                raw = {'source_cells': [v.isoformat() if isinstance(v, (datetime, date)) else v for v in selected],
                       'description': str(selected[2] or ''), 'currency': 'BRL'}
                if any(cells[i].data_type == 'f' for i in headers):
                    errors.append('Fórmula em célula financeira; forneça extrato com valores')
                try:
                    raw.update(trade_date=day(selected[0]), settlement_date=day(selected[1]),
                               amount=str(money(selected[3])), balance=str(money(selected[4])))
                except ValueError as exc:
                    errors.append(str(exc))
                raw['category'] = category(raw['description'])
                symbol = re.search(r'\b([A-Z]{4}\d{1,2})\b', plain(raw['description']))
                raw['symbol'] = symbol[1] if symbol else None
                note = re.search(r'NOTA\s*N[º°O.]?\s*(\d+)', plain(raw['description']))
                raw['note_number'] = note[1] if note else None
                raw['errors'] = errors
                rows.append(SourceRow({'sheet': sheet.title, 'row': number, 'item': len(rows)+1}, raw))
        if not all(metadata.get(k) for k in ('account_number', 'holder', 'period_start', 'period_end')):
            raise ValueError('Cabeçalho XP incompleto: conta, titular ou período ausente')
        if metadata['period_start'] > metadata['period_end']:
            raise ValueError('Período inválido')
        if not rows:
            raise ValueError('Extrato sem movimentos: saldo histórico não pode ser inferido')
        valid = all(not r.values['errors'] for r in rows)
        metadata['errors'] = []
        if valid:
            # Choose chronology by settlement; ties retain documentary order or its reverse.
            options = (rows, list(reversed(rows)))
            def breaks(sequence):
                return sum(Decimal(b.values['balance']) != Decimal(a.values['balance']) + Decimal(b.values['amount'])
                           or b.values['settlement_date'] < a.values['settlement_date'] for a, b in zip(sequence, sequence[1:]))
            chronological = min(options, key=breaks)
            mismatch = breaks(chronological)
            metadata.update(opening_balance=str(Decimal(chronological[0].values['balance']) - Decimal(chronological[0].values['amount'])),
                            closing_balance=chronological[-1].values['balance'],
                            credits=str(sum((Decimal(r.values['amount']) for r in rows if Decimal(r.values['amount']) > 0), Decimal(0))),
                            debits=str(sum((-Decimal(r.values['amount']) for r in rows if Decimal(r.values['amount']) < 0), Decimal(0))),
                            balance_breaks=mismatch, source_locator=chronological[-1].locator)
            if mismatch:
                metadata['errors'].append('Sequência de saldos/datas divergente')
            for r in rows:
                if not metadata['period_start'] <= r.values['settlement_date'] <= metadata['period_end']:
                    r.values['errors'].append('Liquidação fora do período do extrato')
        else:
            metadata['errors'].append('Há linhas financeiras inválidas')
        # Balance and occurrence preserve two legitimate identical movements.
        counts = {}
        for r in rows:
            key = json.dumps([r.values.get(k) for k in ('trade_date','settlement_date','amount','balance')] + [plain(r.values['description'])])
            counts[key] = counts.get(key, 0) + 1
            r.values['fingerprint'] = hashlib.sha256(f'{key}:{counts[key]}'.encode()).hexdigest()
        return metadata, rows
    finally:
        book.close()


class XPStatementAdapter:
    adapter_id = 'xp-account-statement'
    version = '1'
    document_type = 'account_statement'

    def detect(self, filename, body):
        if not body.startswith(b'PK\x03\x04'):
            return 0
        try:
            extract(body)
            return 99
        except ValueError:
            return 0

    def parse(self, filename, body):
        return extract(body)[1]

    def normalize(self, source, db, options=None):
        return dict(source.values, row_number=source.locator['item'], source_locator=source.locator)


XP_ADAPTER = register(XPStatementAdapter())
